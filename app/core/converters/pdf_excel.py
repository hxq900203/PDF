from pathlib import Path

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
import pdfplumber

from ..errors import ConversionError
from ..options import Options
from ..registry import register
from .pdf_image import open_pdf


def clean_cell(value):
    """去掉 Excel 不允许的控制字符。"""
    return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value


@register("pdf", "excel")
def pdf_to_excel(src: Path, out_dir: Path, opts: Options) -> list[Path]:
    open_pdf(src).close()  # 先检查加密/损坏，给出中文提示
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    with pdfplumber.open(src) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            tables = page.extract_tables()
            if not tables:
                continue
            ws = wb.create_sheet(f"第{i}页")
            for t, table in enumerate(tables):
                if t:
                    ws.append([])  # 同页多表之间空一行
                for row in table:
                    ws.append([clean_cell(v) for v in row])
    if not wb.sheetnames:
        raise ConversionError("PDF 中未找到表格")
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{src.stem}.xlsx"
    wb.save(path)
    return [path]
