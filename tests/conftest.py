from pathlib import Path

import docx
import openpyxl
import pymupdf
import pytest

import shutil

NAME = "中文 样例"

needs_soffice = pytest.mark.skipif(shutil.which("soffice") is None, reason="无 LibreOffice")


def _write_text(page, point, text):
    page.insert_text(point, text, fontname="china-s", fontsize=14)


@pytest.fixture
def sample_pdf(tmp_path) -> Path:
    doc = pymupdf.open()
    for i in range(1, 4):
        page = doc.new_page(width=595, height=842)
        _write_text(page, (72, 100), f"第{i}页")
    path = tmp_path / f"{NAME}.pdf"
    doc.save(path)
    return path


TABLE = [["姓名", "分数"], ["张三", "90"], ["李四", "85"]]


@pytest.fixture
def sample_table_pdf(tmp_path) -> Path:
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    x0, y0, cw, rh = 72, 100, 150, 30
    for r in range(len(TABLE) + 1):
        page.draw_line((x0, y0 + r * rh), (x0 + 2 * cw, y0 + r * rh))
    for c in range(3):
        page.draw_line((x0 + c * cw, y0), (x0 + c * cw, y0 + len(TABLE) * rh))
    for r, row in enumerate(TABLE):
        for c, text in enumerate(row):
            _write_text(page, (x0 + c * cw + 8, y0 + r * rh + 20), text)
    path = tmp_path / f"{NAME}.pdf"
    doc.save(path)
    return path


def _png(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 100, 50), False)
    pix.set_rect(pix.irect, (255, 0, 0))
    pix.save(path)
    return path


@pytest.fixture
def sample_png(tmp_path) -> Path:
    return _png(tmp_path / f"{NAME}.png")


@pytest.fixture
def make_png():
    return _png


@pytest.fixture
def sample_docx(tmp_path) -> Path:
    d = docx.Document()
    d.add_paragraph("你好世界")
    path = tmp_path / f"{NAME}.docx"
    d.save(path)
    return path


@pytest.fixture
def sample_xlsx(tmp_path) -> Path:
    wb = openpyxl.Workbook()
    wb.active["A1"], wb.active["B1"] = "名称", "数量"
    path = tmp_path / f"{NAME}.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def encrypted_pdf(tmp_path) -> Path:
    doc = pymupdf.open()
    doc.new_page()
    path = tmp_path / "加密.pdf"
    doc.save(path, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="1", owner_pw="2")
    return path


@pytest.fixture
def corrupt_pdf(tmp_path) -> Path:
    path = tmp_path / "损坏.pdf"
    path.write_bytes(b"not a pdf")
    return path
